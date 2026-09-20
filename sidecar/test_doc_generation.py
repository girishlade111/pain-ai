"""pain ai — Pytest Suite for Document Generation (Phase 10)

Drives the SAME Hermes skill helper scripts the agent invokes via the
terminal tool (no mocks): one real PDF, DOCX, XLSX, and PPTX each, plus a
source-code project scaffold. Every file is verified: exists, non-zero,
correct extension, reopens/parses cleanly (no corruption).
"""

import json
import subprocess
import sys
from pathlib import Path

root_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(root_dir / "sidecar"))
sys.path.insert(0, str(root_dir / "hermes-agent"))

SKILLS = root_dir / "hermes-agent" / "skills" / "productivity"

PDF_SCRIPT = SKILLS / "pdf" / "scripts" / "pdf_create.py"
DOCX_SCRIPT = SKILLS / "docx" / "scripts" / "docx_create.py"
XLSX_SCRIPT = SKILLS / "xlsx" / "scripts" / "xlsx_create.py"
PPTX_SCRIPT = SKILLS / "powerpoint" / "scripts" / "pptx_create.py"


def _run(script: Path, *args: str) -> None:
    """Run a skill helper exactly as the agent would (subprocess, JSON CLI)."""
    assert script.is_file(), f"skill script missing: {script}"
    proc = subprocess.run(
        [sys.executable, str(script), *args],
        capture_output=True, text=True, timeout=180)
    assert proc.returncode == 0, (
        f"{script.name} failed (exit {proc.returncode}):\n{proc.stderr[-2000:]}")


def _write_spec(tmp_path: Path, name: str, spec: dict) -> Path:
    path = tmp_path / name
    path.write_text(json.dumps(spec), encoding="utf-8")
    return path


def _assert_healthy(path: Path, suffix: str) -> None:
    assert path.is_file(), f"not generated: {path}"
    assert path.suffix.lower() == suffix, path.suffix
    assert path.stat().st_size > 0, f"zero-size file: {path}"


def test_generate_pdf_and_reopen(tmp_path):
    out = tmp_path / "report.pdf"
    spec = _write_spec(tmp_path, "report-spec.json", {
        "title": "pain ai Phase 10 Report",
        "author": "pain ai",
        "elements": [
            {"type": "heading", "text": "Generation Proof", "level": 1},
            {"type": "paragraph", "text": "Real PDF produced by the Hermes pdf skill."},
            {"type": "table", "rows": [["Metric", "Value"], ["tests", "green"]], "header": True},
        ],
    })
    _run(PDF_SCRIPT, str(spec), "-o", str(out))
    _assert_healthy(out, ".pdf")

    from pypdf import PdfReader
    reader = PdfReader(str(out))
    assert len(reader.pages) >= 1, "PDF has no pages (corrupt)"
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    assert "Generation Proof" in text, "expected heading missing on reopen"


def test_generate_docx_and_reopen(tmp_path):
    out = tmp_path / "memo.docx"
    spec = _write_spec(tmp_path, "memo-spec.json", {
        "blocks": [
            {"type": "heading", "text": "Phase 10 Memo", "level": 1},
            {"type": "paragraph", "text": "Real DOCX produced by the Hermes docx skill."},
            {"type": "bullet_list", "items": ["first item", "second item"]},
            {"type": "table", "header": ["Col1", "Col2"],
             "rows": [["a", "b"], ["c", "d"]]},
        ],
    })
    _run(DOCX_SCRIPT, str(spec), str(out))
    _assert_healthy(out, ".docx")

    from docx import Document
    doc = Document(str(out))
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "Phase 10 Memo" in full, "expected heading missing on reopen"
    assert len(doc.tables) >= 1, "expected table missing on reopen"


def test_generate_xlsx_and_reopen(tmp_path):
    out = tmp_path / "budget.xlsx"
    spec = _write_spec(tmp_path, "budget-spec.json", {
        "sheets": [
            {"name": "Q1",
             "rows": [["Item", "Cost"], ["Widgets", 120], ["Gadgets", 75]]},
            {"name": "Summary",
             "cells": {"A1": {"value": "Total"}, "B1": {"value": 195}}},
        ],
    })
    _run(XLSX_SCRIPT, str(spec), str(out))
    _assert_healthy(out, ".xlsx")

    from openpyxl import load_workbook
    wb = load_workbook(str(out), data_only=False)
    assert set(wb.sheetnames) >= {"Q1", "Summary"}, wb.sheetnames
    ws = wb["Q1"]
    assert ws["A1"].value == "Item" and ws["B2"].value == 120
    assert wb["Summary"]["B1"].value == 195


def test_generate_pptx_and_reopen(tmp_path):
    out = tmp_path / "deck.pptx"
    spec = _write_spec(tmp_path, "deck-spec.json", {
        "slides": [
            {"layout": "title", "title": "Phase 10 Deck", "subtitle": "Hermes powerpoint skill"},
            {"layout": "title_content", "title": "Agenda",
             "bullets": ["Generate", "Verify", "Ship"]},
        ],
    })
    _run(PPTX_SCRIPT, str(spec), str(out))
    _assert_healthy(out, ".pptx")

    from pptx import Presentation
    deck = Presentation(str(out))
    assert len(deck.slides) == 2, "expected 2 slides on reopen"
    texts = [shape.text for slide in deck.slides for shape in slide.shapes
             if shape.has_text_frame]
    assert any("Phase 10 Deck" in t for t in texts)
    assert any("Verify" in t for t in texts)


def test_source_code_project_scaffold_and_tracking(tmp_path):
    """Source-code projects: scaffold files, verify structure, and prove the
    artifact pipeline (Phases 5/6) tracks code projects file-by-file."""
    from output_manager import diff_artifacts, snapshot_dir
    from artifact_store import build_group

    project = tmp_path / "demo-app"
    (project / "src").mkdir(parents=True)
    before = snapshot_dir(tmp_path)
    (project / "package.json").write_text(
        '{"name": "demo-app", "version": "1.0.0"}', encoding="utf-8")
    (project / "src" / "index.ts").write_text(
        'export function main(): void {\n  console.log("hi");\n}\n',
        encoding="utf-8")
    (project / "README.md").write_text("# demo-app\n", encoding="utf-8")
    (project / "index.html").write_text("<!doctype html><html></html>\n",
                                        encoding="utf-8")

    detected = diff_artifacts(tmp_path, before)
    assert len(detected) == 4, [a["path"] for a in detected]
    group = build_group(tmp_path,
                        [a["path"] for a in detected], "sess-1", "turn-1")
    assert group["kind"] == "project" and group["status"] == "completed"
    kinds = {r["filename"]: r["type"] for r in group["files"]}
    assert kinds["package.json"] == "code"
    assert kinds["index.ts"] == "code"
    assert all(r["size"] > 0 for r in group["files"])


def test_skill_toolchain_presence():
    """Discovery contract: the four productivity skills + create scripts the
    agent invokes must exist (execution path components, not the LLM)."""
    for skill, script in (("pdf", PDF_SCRIPT), ("docx", DOCX_SCRIPT),
                          ("xlsx", XLSX_SCRIPT), ("powerpoint", PPTX_SCRIPT)):
        assert (SKILLS / skill / "SKILL.md").is_file()
        assert script.is_file(), f"missing create script for {skill}"


def test_agent_file_write_path_components():
    """Source-code generation reaches disk through Hermes write_file, gated
    by the desktop boundary as FileWrite (no parallel executor)."""
    from tools.registry import discover_builtin_tools, registry
    from capability_gate import tool_action

    discover_builtin_tools()
    assert registry.get_entry("write_file") is not None
    action = tool_action("write_file", {"path": "src/app.ts"})
    assert action is not None and action["kind"] == "FileWrite"
